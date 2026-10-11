"""P7 多格式文档解析器单元测试。"""

from __future__ import annotations

import os
import tempfile

import pytest

from app.services.indexing.document_parser import parse_document


class TestParseDocument:
    def test_txt_utf8(self) -> None:
        with tempfile.NamedTemporaryFile(
            mode="w", suffix=".txt", delete=False, encoding="utf-8"
        ) as fh:
            fh.write("第一行\n第二行\n第三行")
            path = fh.name
        try:
            result = parse_document(path)
            assert result.format == "txt"
            assert result.char_count > 0
            assert "第一行" in result.text
            assert "第三行" in result.text
        finally:
            os.unlink(path)

    def test_txt_gbk_fallback(self) -> None:
        with tempfile.NamedTemporaryFile(suffix=".txt", delete=False) as fh:
            fh.write("中文内容 GBK 编码".encode("gbk"))
            path = fh.name
        try:
            result = parse_document(path)
            assert result.format == "txt"
            assert "中文" in result.text
        finally:
            os.unlink(path)

    def test_markdown(self) -> None:
        with tempfile.NamedTemporaryFile(
            mode="w", suffix=".md", delete=False, encoding="utf-8"
        ) as fh:
            fh.write("# 标题\n\n正文段落。")
            path = fh.name
        try:
            result = parse_document(path)
            assert result.format == "md"
            assert "标题" in result.text
            assert "正文" in result.text
        finally:
            os.unlink(path)

    def test_unsupported_format(self) -> None:
        with tempfile.NamedTemporaryFile(suffix=".xyz", delete=False) as fh:
            fh.write(b"dummy")
            path = fh.name
        try:
            with pytest.raises(ValueError, match="Unsupported document format"):
                parse_document(path)
        finally:
            os.unlink(path)

    def test_parsed_document_dataclass(self) -> None:
        with tempfile.NamedTemporaryFile(
            mode="w", suffix=".txt", delete=False, encoding="utf-8"
        ) as fh:
            fh.write("A" * 42)
            path = fh.name
        try:
            result = parse_document(path)
            assert result.char_count == 42
            assert result.page_count == 1  # 非 PDF 的默认值
        finally:
            os.unlink(path)

    def test_pdf_minimal(self) -> None:
        """解析一个内存中构造的极小 PDF。需要 PyMuPDF。"""
        try:
            import pymupdf  # noqa: F401
        except ImportError:
            pytest.skip("PyMuPDF not installed")

        # 在内存中构造一个极小的单页 PDF。
        import pymupdf

        with tempfile.NamedTemporaryFile(
            suffix=".pdf", delete=False
        ) as fh:
            path = fh.name
        try:
            doc = pymupdf.open()
            page = doc.new_page()
            page.insert_text((50, 100), "动火作业安全规定 PDF 测试")
            doc.save(path)
            doc.close()

            result = parse_document(path)
            assert result.format == "pdf"
            assert "动火" in result.text or "PDF" in result.text
        finally:
            os.unlink(path)

    def test_docx_minimal(self) -> None:
        """解析一个内存中构造的极小 DOCX。需要 python-docx。"""
        try:
            import docx  # noqa: F401
        except ImportError:
            pytest.skip("python-docx not installed")

        import docx

        with tempfile.NamedTemporaryFile(
            suffix=".docx", delete=False
        ) as fh:
            path = fh.name
        try:
            d = docx.Document()
            d.add_paragraph("动火作业审批流程 DOCX 测试")
            # 添加一个小表格——解析器也应能识别。
            t = d.add_table(rows=1, cols=2)
            t.rows[0].cells[0].text = "审批人"
            t.rows[0].cells[1].text = "安全员"
            d.save(path)

            result = parse_document(path)
            assert result.format == "docx"
            assert "动火" in result.text
            assert "审批人" in result.text
            assert "安全员" in result.text
        finally:
            os.unlink(path)
