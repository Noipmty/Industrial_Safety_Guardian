"""多格式文档解析器：PDF / DOCX / TXT / Markdown。

每个解析器返回纯文本字符串，供下游 ``recursive_chunk`` 阶段使用。
PDF 使用 PyMuPDF（``import pymupdf``），DOCX 使用 python-docx，
TXT / Markdown 则退化为按编码直接读文件。
"""

from __future__ import annotations

import os
from collections.abc import Callable
from dataclasses import dataclass

from app.core.logger import get_logger

logger = get_logger(__name__)


@dataclass(slots=True)
class ParsedDocument:
    """解析文档文件的结果。"""

    text: str
    format: str  # 取值："pdf" | "docx" | "txt" | "md"
    page_count: int = 1
    char_count: int = 0

    def __post_init__(self) -> None:
        self.char_count = len(self.text)


def _parse_txt(path: str) -> str:
    """读取纯文本文件，utf-8 失败时回退 gbk。"""
    try:
        with open(path, encoding="utf-8") as fh:
            return fh.read()
    except UnicodeDecodeError:
        with open(path, encoding="gbk") as fh:
            return fh.read()


def _parse_pdf(path: str) -> str:
    """通过 PyMuPDF 提取 PDF 文本。"""
    import pymupdf  # noqa: PLC0415

    doc = pymupdf.open(path)
    try:
        pages: list[str] = []
        for i in range(doc.page_count):
            pages.append(doc.load_page(i).get_text("text"))
        text = "\n\n".join(pages)
        # 清理过多的空行。
        while "\n\n\n" in text:
            text = text.replace("\n\n\n", "\n\n")
        return text.strip()
    finally:
        doc.close()


def _parse_docx(path: str) -> str:
    """通过 python-docx 提取 DOCX 文本。"""
    import docx  # 惰性导入  # noqa: PLC0415

    d = docx.Document(path)
    paragraphs = [p.text for p in d.paragraphs if p.text.strip()]
    # 同时读取表格单元格 —— 其中常包含结构化知识。
    for table in d.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                paragraphs.append(" | ".join(cells))
    return "\n\n".join(paragraphs)


# 扩展名 → (解析函数, 格式名)
_PARSERS: dict[str, tuple[Callable[[str], str], str]] = {
    ".pdf": (_parse_pdf, "pdf"),
    ".docx": (_parse_docx, "docx"),
    ".txt": (_parse_txt, "txt"),
    ".md": (_parse_txt, "md"),
    ".markdown": (_parse_txt, "md"),
}


def parse_document(path: str) -> ParsedDocument:
    """解析任意受支持的文档，返回纯文本 + 元数据。

    不支持的扩展名会抛出 ``ValueError``。
    """
    ext = os.path.splitext(path)[1].lower()
    if ext not in _PARSERS:
        supported = ", ".join(sorted(_PARSERS))
        raise ValueError(
            f"Unsupported document format '{ext}'. Supported: {supported}"
        )

    parser, fmt = _PARSERS[ext]
    logger.info("Parsing %s document: %s", fmt, path)
    text = parser(path)
    return ParsedDocument(text=text, format=fmt)
